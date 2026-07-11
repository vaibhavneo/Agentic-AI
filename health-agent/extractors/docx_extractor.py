from __future__ import annotations
"""
DOCX extractor: reads health reports, doctor notes, clinical summaries in Word format.
"""
import json
import re
import anthropic
from docx import Document

_client = anthropic.Anthropic()

EXTRACTION_PROMPT = """You are a medical data extractor. Parse this health document and extract ALL vital signs.

Document text:
{text}

Return ONLY valid JSON (no markdown, no extra text):
{{"patient_id": "<string or UNKNOWN>", "timestamp": "<ISO string or null>", "heart_rate": <bpm or null>, "systolic_bp": <mmHg or null>, "diastolic_bp": <mmHg or null>, "spo2": <percentage or null>, "temperature": <Celsius or null>, "respiratory_rate": <per min or null>, "glucose": <mg/dL or null>, "weight": <kg or null>, "notes": "<clinical notes or null>"}}

Convert units if needed (F to C, mmol/L to mg/dL)."""


def extract_from_docx(docx_path: str) -> dict:
    """Extract vitals from a Word document health report."""
    doc = Document(docx_path)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_cells = []
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                table_cells.append(" | ".join(row_texts))

    full_text = "\n".join(paragraphs + table_cells)
    if not full_text.strip():
        raise ValueError("No extractable text found in DOCX: " + docx_path)

    if len(full_text) > 12000:
        full_text = full_text[:12000] + "\n[... truncated ...]"

    response = _client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=1024,
        messages=[{"role": "user", "content": EXTRACTION_PROMPT.format(text=full_text)}],
    )

    text = response.content[0].text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-z]*\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return json.loads(text)
